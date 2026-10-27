"""Stage 2: Map. Build the cell dependency graph, find outputs and true inputs.

Writes graph.json (machine) and map.md (human). Dynamic references (OFFSET, INDIRECT)
become edges with confidence "dynamic" plus the candidate cells they could touch.
"""
import datetime as dt
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

import networkx as nx
import openpyxl
from openpyxl.formula.tokenizer import Token
from openpyxl.utils import get_column_letter

from . import refs as R
from .excelfn import _rank
from .inventory import formula_text
from .xlsxml import Package

OUTPUT_SHEET = re.compile(r"(?i)summary|output|dashboard|report|results|overview")
AGGREGATES = {"SUM", "SUMIF", "SUMIFS", "SUMPRODUCT", "AVERAGE", "AVERAGEIF", "AVERAGEIFS", "COUNT", "COUNTA",
              "COUNTIF", "COUNTIFS", "MAX", "MIN", "SUBTOTAL", "AGGREGATE"}
LOOKUPS = {"VLOOKUP", "HLOOKUP", "MATCH", "LOOKUP", "XLOOKUP", "XMATCH"}


def jsonable(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return {"date": v.isoformat()}
    if isinstance(v, (int, float, str, bool)) or v is None:
        return v
    return str(v)


class Book:
    """Formula + cached-value view of a workbook with name resolution."""

    def __init__(self, path):
        self.path = Path(path)
        self.f = openpyxl.load_workbook(path, keep_vba=True)
        self.v = openpyxl.load_workbook(path, keep_vba=True, data_only=True)
        info = Package(path).workbook()
        self.states = {s["name"]: s["state"] for s in info["sheets"]}
        self.names = {}
        for dn in info["defined_names"]:
            self.names[(dn["scope"], dn["name"].lower())] = dn["refers_to"]
        self.dims = {ws.title: (ws.max_row, ws.max_column) for ws in self.f.worksheets}

    def resolve_name(self, name, sheet):
        txt = self.names.get((sheet, name.lower())) or self.names.get((None, name.lower()))
        if txt is None:
            return None
        return [R.parse_ref(part, sheet) for part in txt.split(",")]

    def value(self, key):
        sheet, c, r = R.split_key(key)
        if sheet not in self.dims:
            return None
        return self.v[sheet].cell(r, c).value

    def formula(self, key):
        sheet, c, r = R.split_key(key)
        if sheet not in self.dims:
            return None
        return formula_text(self.f[sheet].cell(r, c).value)

    def expand(self, ref):
        if ref.sheet not in self.dims:
            return []
        mr, mc = self.dims[ref.sheet]
        return ref.cells(max_row=mr, max_col=mc)


def build(path, out_dir):
    out_dir = Path(out_dir)
    bk = Book(path)
    G = nx.DiGraph()
    dynamic, broken, lookups, gaps = [], [], [], []
    thresholds = defaultdict(set)          # formula cell -> comparison literals

    for ws in bk.f.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                key = f"{ws.title}!{cell.coordinate}"
                f = formula_text(cell.value)
                G.add_node(key, sheet=ws.title, formula=f, kind="formula" if f else "constant")
                if not f:
                    continue
                for op, lit in R.comparisons(f):
                    thresholds[key].add(lit)
                _add_edges(G, bk, key, ws.title, f, dynamic, broken)
                _lookup_checks(bk, key, ws.title, f, lookups)
                _aggregate_gaps(bk, key, ws.title, f, gaps)

    for n in G.nodes:
        if "kind" not in G.nodes[n]:
            G.nodes[n].update(sheet=n.rsplit("!", 1)[0], formula=None,
                              kind="external" if n.startswith("[") else "blank")

    formulas = [n for n, d in G.nodes(data=True) if d["kind"] == "formula"]
    out_sheets = [s for s in bk.dims if OUTPUT_SHEET.search(s)]
    sinks = [n for n in formulas if G.out_degree(n) == 0]
    if out_sheets:
        outputs = [n for n in formulas if G.nodes[n]["sheet"] in out_sheets]
        secondary = [n for n in sinks if G.nodes[n]["sheet"] not in out_sheets]
        rule = f"formula cells on sheet(s) {out_sheets}"
    else:
        outputs, secondary = sinks, []
        rule = "formula cells nothing else reads (no Summary/Output sheet found)"
    outputs = sorted(outputs, key=_sort_key)

    ancestors = set()
    for o in outputs:
        ancestors |= nx.ancestors(G, o)
    inputs = sorted((n for n in ancestors if G.nodes[n]["kind"] == "constant"), key=_sort_key)
    externals = sorted(n for n in ancestors if G.nodes[n]["kind"] == "external")

    # thresholds that sit downstream of each input (within 4 hops)
    in_thresh = {}
    for n in inputs:
        lits = set()
        for d in nx.single_source_shortest_path_length(G, n, cutoff=4):
            lits |= thresholds.get(d, set())
        in_thresh[n] = sorted(lits)
    blocks = _input_blocks(bk, inputs, in_thresh)

    cycles = [sorted(c, key=_sort_key) for c in nx.strongly_connected_components(G) if len(c) > 1]
    cycles += [[n] for n in G.nodes if G.has_edge(n, n)]

    graph = dict(
        file=str(path), output_rule=rule, outputs=outputs, outputs_secondary=sorted(secondary, key=_sort_key),
        inputs=inputs, input_blocks=blocks, externals=externals,
        dynamic=dynamic, broken_refs=broken, lookups=_dedupe(lookups), aggregate_gaps=_dedupe(gaps),
        cycles=[dict(size=len(c), cells=c[:60]) for c in cycles],
        stats=dict(nodes=G.number_of_nodes(), edges=G.number_of_edges(), formulas=len(formulas),
                   outputs=len(outputs), inputs=len(inputs), dynamic_edges=sum(len(d["candidates"]) for d in dynamic)),
        nodes={n: dict(sheet=d["sheet"], kind=d["kind"], formula=d["formula"], value=jsonable(bk.value(n))
                       if not n.startswith("[") else None) for n, d in G.nodes(data=True)},
        edges=[[u, v, d.get("confidence", "static")] for u, v, d in G.edges(data=True)],
    )
    (out_dir / "graph.json").write_text(json.dumps(graph, default=str))
    (out_dir / "map.md").write_text(render_map(bk, G, graph))
    return graph


def _sort_key(key):
    if key.startswith("["):
        return (key, 0, 0)
    s, c, r = R.split_key(key)
    return (s, c, r)


def _dedupe(items, ignore=("cells", "count", "cell", "formula", "sheet")):
    """Merge items that differ only by which cell they came from (copied-down formulas)."""
    merged = {}
    for it in items:
        k = json.dumps({k: v for k, v in it.items() if k not in ignore}, sort_keys=True, default=str)
        if k in merged:
            merged[k]["count"] += 1
            if len(merged[k]["cells"]) < 8:
                merged[k]["cells"].extend(it.get("cells", []))
        else:
            merged[k] = dict(it, count=1, cells=list(it.get("cells", [])))
    return list(merged.values())


# ---------------------------------------------------------------- edges

def _add_edges(G, bk, key, sheet, f, dynamic, broken):
    node_tree = R.tree(f)
    for r in R.refs(f, sheet):
        targets = []
        if r.kind == "name":
            resolved = bk.resolve_name(r.text, sheet)
            if resolved is None:
                broken.append(dict(cell=key, ref=r.text, reason="unknown name"))
                continue
            for rr in resolved:
                if rr.kind == "error":
                    broken.append(dict(cell=key, ref=r.text, reason=f"name refers to {rr.text}"))
                else:
                    targets += bk.expand(rr)
        elif r.kind == "error":
            broken.append(dict(cell=key, ref=r.text, reason="#REF! in formula"))
            continue
        elif r.kind == "external":
            targets = [f"[{r.external}]{r.sheet}!{get_column_letter(r.c1)}{r.r1}"]
        else:
            try:
                targets = bk.expand(r)
            except ValueError:
                broken.append(dict(cell=key, ref=r.text, reason="range too large to expand"))
                continue
        conf = "dynamic-base" if (r.func == "OFFSET" and r.arg == 0) else "static"
        for t in targets:
            G.add_edge(t, key, confidence=conf)

    for n, func, arg, sib, i in R.walk(node_tree):
        if isinstance(n, R.Func) and n.name == "OFFSET":
            dynamic.append(_offset_candidates(G, bk, key, sheet, n))
        elif isinstance(n, R.Func) and n.name == "INDIRECT":
            dynamic.append(_indirect_candidates(G, bk, key, sheet, n))


def _refs_in(nodes, sheet):
    return R.refs("=" + R.node_text(nodes), sheet) if nodes else []


def _offset_candidates(G, bk, key, sheet, fn):
    base = _refs_in(fn.args[0], sheet) if fn.args else []
    info = dict(cell=key, function="OFFSET", formula=bk.formula(key), base=[b.a1 for b in base],
                drivers=[], candidates=[])
    if not base or base[0].kind not in ("cell", "range"):
        return info
    b = base[0]
    width, height = 1, 1
    for idx, dim in ((1, "rows"), (2, "cols")):
        if idx < len(fn.args):
            for r in _refs_in(fn.args[idx], sheet):
                info["drivers"].append(r.a1)
                if r.kind == "range" and r.r1 == r.r2 and r.c1 and r.c2:
                    width = max(width, r.c2 - r.c1 + 1)
                if r.kind == "range" and r.c1 == r.c2 and r.r1 and r.r2:
                    height = max(height, r.r2 - r.r1 + 1)
    from .refs import Ref
    cand = Ref("cand", "range", b.sheet, b.c1, b.r1, (b.c2 or b.c1) + width - 1, (b.r2 or b.r1) + height - 1)
    cells = bk.expand(cand)
    for c in cells:
        if not G.has_edge(c, key):
            G.add_edge(c, key, confidence="dynamic")
    info["candidates"] = [cand.a1]
    info["candidate_cells"] = len(cells)
    return info


def _indirect_candidates(G, bk, key, sheet, fn):
    arg = fn.args[0] if fn.args else []
    text = R.node_text(arg)
    info = dict(cell=key, function="INDIRECT", formula=bk.formula(key), argument=text, drivers=[], candidates=[])
    lits = [t.value.strip('"') for t in arg if isinstance(t, Token) and t.subtype == Token.TEXT]
    for r in _refs_in(arg, sheet):
        info["drivers"].append(r.a1)
    joined = "".join(lits)
    m = re.search(r"!(\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)", joined)
    if len(lits) == 1 and not info["drivers"]:
        r = R.parse_ref(lits[0], sheet)
        cells = bk.expand(r) if r.kind in ("cell", "range") else []
        for c in cells:
            G.add_edge(c, key, confidence="indirect-static")
        info["candidates"] = [r.a1]
    elif m:
        addr = m.group(1).replace("$", "")
        for s in bk.dims:
            r = R.parse_ref(f"'{s}'!{addr}", s)
            for c in bk.expand(r):
                if bk.value(c) is not None or bk.formula(c):
                    G.add_edge(c, key, confidence="dynamic")
            info["candidates"].append(f"'{s}'!{addr}")
        info["note"] = f"sheet chosen at run time; address {addr} on any of {len(bk.dims)} sheets"
    else:
        info["note"] = "target built at run time; not resolvable statically"
    return info


# ---------------------------------------------------------------- checks

def _arg_is_true(nodes, default=True):
    if not nodes:
        return default
    t = R.node_text(nodes).strip().upper()
    if t in ("TRUE", "1"):
        return True
    if t in ("FALSE", "0"):
        return False
    return None                                   # computed at run time


def _lookup_checks(bk, key, sheet, f, out):
    for fn in R.functions(f):
        if fn.name not in LOOKUPS:
            continue
        table_arg = {"VLOOKUP": 1, "HLOOKUP": 1, "MATCH": 1, "LOOKUP": 1, "XLOOKUP": 1, "XMATCH": 1}[fn.name]
        if len(fn.args) <= table_arg:
            continue
        tr = _refs_in(fn.args[table_arg], sheet)
        if not tr or tr[0].kind not in ("cell", "range", "name"):
            continue
        t = tr[0]
        if t.kind == "name":
            res = bk.resolve_name(t.text, sheet)
            if not res or res[0].kind not in ("cell", "range"):
                continue
            t = res[0]
        if fn.name in ("VLOOKUP", "LOOKUP", "MATCH") or fn.name == "HLOOKUP":
            pos = {"VLOOKUP": 3, "HLOOKUP": 3, "MATCH": 2}.get(fn.name)
            if fn.name == "LOOKUP":
                approx = True
            elif fn.name == "MATCH":
                mt = R.node_text(fn.args[2]).strip() if len(fn.args) > 2 else "1"
                approx = True if mt in ("1", "") else False if mt == "0" else None
            else:
                approx = _arg_is_true(fn.args[pos] if len(fn.args) > pos else None)
        else:
            approx = False
        if fn.name == "HLOOKUP":
            keys_ref = R.Ref("k", "range", t.sheet, t.c1, t.r1, t.c2, t.r1)
        elif fn.name in ("MATCH", "XMATCH", "XLOOKUP", "LOOKUP"):
            keys_ref = t
        else:
            keys_ref = R.Ref("k", "range", t.sheet, t.c1, t.r1, t.c1, t.r2)
        cells = bk.expand(keys_ref)
        keys = [bk.value(c) for c in cells]
        while keys and keys[-1] is None:
            keys.pop()
        item = dict(function=fn.name, table=t.a1, keys=keys_ref.a1, approximate=approx, cells=[key])
        problems = []
        nonblank = [k for k in keys if k is not None]
        if approx:
            try:
                ranks = [_rank(k) for k in nonblank]
                if any(a > b for a, b in zip(ranks, ranks[1:])):
                    firsts = [i for i, (a, b) in enumerate(zip(ranks, ranks[1:])) if a > b][:3]
                    problems.append(f"approximate match on keys that are NOT sorted ascending "
                                    f"(first break after position {firsts[0] + 1}: {nonblank[firsts[0]]!r} > {nonblank[firsts[0] + 1]!r})")
            except Exception:
                problems.append("approximate match on keys of mixed types")
        norm = [k.strip().lower() if isinstance(k, str) else k for k in nonblank]
        dupes = [k for k, c in Counter(norm).items() if c > 1]
        if dupes:
            problems.append(f"duplicate keys {dupes[:5]} (first match wins)")
        ws_keys = [k for k in nonblank if isinstance(k, str) and k != k.strip()]
        if ws_keys:
            problems.append(f"keys with leading/trailing spaces {ws_keys[:5]}")
        kinds = Counter(type(k).__name__ for k in nonblank)
        if len(kinds) > 1:
            problems.append(f"key column mixes types {dict(kinds)}")
        if approx is None:
            problems.append("match mode computed at run time")
        item["problems"] = problems
        item["key_count"] = len(nonblank)
        out.append(item)


def _aggregate_gaps(bk, key, sheet, f, out):
    s0, c0, r0 = R.split_key(key)
    for r in R.refs(f, sheet):
        if r.func not in AGGREGATES or r.kind != "range" or r.whole_col or r.whole_row:
            continue
        if r.c1 != r.c2 or r.r2 - r.r1 < 3:
            continue
        below = []
        rr = r.r2 + 1
        last_key = R.cell_key(r.sheet, r.c1, r.r2)
        last_pat = R.r1c1(bk.formula(last_key), r.r2, r.c1) if bk.formula(last_key) else None
        while rr <= bk.dims.get(r.sheet, (0, 0))[0] and len(below) < 200:
            k = R.cell_key(r.sheet, r.c1, rr)
            v, fx = bk.value(k), bk.formula(k)
            if k == key:
                break
            if fx:                                   # a continuing copy-down of the same formula counts as data
                if last_pat is None or R.r1c1(fx, rr, r.c1) != last_pat:
                    break
            elif not isinstance(v, (int, float)) or isinstance(v, bool):
                break
            below.append(k)
            rr += 1
        if below:
            L = get_column_letter(r.c1)
            out.append(dict(pattern=R.r1c1(bk.formula(key), r0, c0), formula=bk.formula(key),
                            range=f"{L}{r.r1}:{L}{r.r2}", excluded_after=f"{L}{r.r2 + 1}:{L}{r.r2 + len(below)}",
                            excluded_count=len(below), cells=[key],
                            note="numeric data continues directly below the aggregated range"))


# ---------------------------------------------------------------- inputs

def _vtype(v):
    if isinstance(v, bool):
        return "bool"
    if isinstance(v, (int, float)):
        return "number"
    if isinstance(v, (dt.date, dt.datetime)):
        return "date"
    return "text"


def _input_blocks(bk, inputs, in_thresh):
    by_col = defaultdict(list)
    for n in inputs:
        s, c, r = R.split_key(n)
        by_col[(s, c)].append(r)
    blocks = []
    for (s, c), rows in sorted(by_col.items()):
        rows.sort()
        run = [rows[0]]
        for r in rows[1:] + [None]:
            if r is not None and r == run[-1] + 1:
                run.append(r)
                continue
            cells = [R.cell_key(s, c, x) for x in run]
            vals = [bk.value(k) for k in cells]
            types = Counter(_vtype(v) for v in vals if v is not None)
            nums = [v for v in vals if _vtype(v) == "number"]
            dates = [v for v in vals if _vtype(v) == "date"]
            texts = [v for v in vals if _vtype(v) == "text"]
            letter = get_column_letter(c)
            header = None
            for hr in range(run[0] - 1, 0, -1):
                hv = bk.v[s].cell(hr, c).value
                if isinstance(hv, str):
                    header = hv
                    break
            blk = dict(range=f"{s}!{letter}{run[0]}" + (f":{letter}{run[-1]}" if len(run) > 1 else ""),
                       sheet=s, col=letter, first_row=run[0], last_row=run[-1], cells=len(run),
                       header=header, types=dict(types), sheet_state=bk.states.get(s, "visible"),
                       thresholds=sorted({t for k in cells for t in in_thresh.get(k, [])}))
            if nums:
                blk.update(min=min(nums), max=max(nums), integers=all(float(x).is_integer() for x in nums))
            if dates:
                blk.update(min_date=min(dates).isoformat(), max_date=max(dates).isoformat())
            if texts:
                distinct = sorted(set(texts), key=str)
                blk.update(distinct=distinct[:30], distinct_count=len(distinct))
            blocks.append(blk)
            if r is not None:
                run = [r]
    return blocks




# ---------------------------------------------------------------- map.md

def _blocks_of(bk, G):
    """One block per contiguous run of connected cells in a column (formulas and typed values together).
    Unreferenced cells (headers, notes) are left out."""
    by_col = defaultdict(list)
    for n, d in G.nodes(data=True):
        if n.startswith("[") or d["kind"] == "blank" or G.degree(n) == 0:
            continue
        s, c, r = R.split_key(n)
        by_col[(s, c)].append((r, n))
    cell_block, blocks = {}, {}
    for (s, c), items in by_col.items():
        items.sort()
        cur = None
        for r, n in items:
            if not (cur and r == cur["last"] + 1):
                cur = dict(sheet=s, col=c, first=r, last=r, cells=[])
                blocks[id(cur)] = cur
            cur["last"] = r
            cur["cells"].append(n)
            cell_block[n] = id(cur)
    for b in blocks.values():
        L = get_column_letter(b["col"])
        b["name"] = f"{b['sheet']}!{L}{b['first']}" + (f":{L}{b['last']}" if b["last"] > b["first"] else "")
        pats = defaultdict(list)
        consts = []
        for n in b["cells"]:
            f = bk.formula(n)
            if f:
                _, c, r = R.split_key(n)
                pats[R.r1c1(f, r, c)].append(n)
            else:
                consts.append(n)
        b["patterns"] = sorted(pats.values(), key=len, reverse=True)
        b["constants"] = consts
    return cell_block, blocks


def _rows(cells, limit=8):
    rows = [R.split_key(c)[2] for c in cells]
    return ", ".join(map(str, rows[:limit])) + (f" ... (+{len(rows) - limit})" if len(rows) > limit else "")


def _label(bk, blocks, bid):
    if bid not in blocks:
        return f"{bid} (external workbook)" if str(bid).startswith("[") else str(bid)
    b = blocks[bid]
    state = bk.states.get(b["sheet"], "visible")
    tag = f" [{state} sheet]" if state != "visible" else ""
    if not b["patterns"]:
        vals = [bk.value(c) for c in b["constants"][:3]]
        return f"{b['name']} typed values ({len(b['constants'])}) e.g. {vals}{tag}"
    main = b["patterns"][0]
    s = f"{b['name']} {bk.formula(main[0])}" + (f"  (x{len(main)})" if len(b["cells"]) > 1 else "") + tag
    for p in b["patterns"][1:3]:
        s += f"\n      ! {len(p)} cell(s) use a different formula, rows {_rows(p)}: {bk.formula(p[0])}"
    if len(b["patterns"]) > 3:
        s += f"\n      ! ...and {len(b['patterns']) - 3} more formula variants"
    if b["constants"]:
        s += f"\n      ! {len(b['constants'])} typed value(s) inside this formula run, rows {_rows(b['constants'])}"
    return s


def _fmt(item, keys):
    return "; ".join(f"{k}: {item[k]}" for k in keys if item.get(k) not in (None, [], ""))


def render_map(bk, G, graph):
    cell_block, blocks = _blocks_of(bk, G)
    bedges, conf_of = defaultdict(set), {}
    for u, v, d in G.edges(data=True):
        bu, bv = cell_block.get(u, u), cell_block.get(v, v)
        if bu != bv:
            bedges[bv].add(bu)
            if d.get("confidence", "static") != "static":
                conf_of[(bu, bv)] = d["confidence"]
    label = lambda bid: _label(bk, blocks, bid)
    short = lambda bid: blocks[bid]["name"] if bid in blocks else str(bid)

    st = graph["stats"]
    lines = [f"# Map: {Path(graph['file']).name}", "",
             f"Outputs = {graph['output_rule']}. {st['outputs']} output cells, {st['inputs']} true input cells "
             f"in {len(graph['input_blocks'])} blocks, {st['formulas']} formulas, {len(graph['dynamic'])} dynamic "
             f"references, {len(graph['cycles'])} circular group(s).", "",
             "Each line is a column block. `!` lines are cells that break the block's pattern. "
             "`<~ dynamic` is an edge the static map cannot confirm (OFFSET/INDIRECT).", "",
             "## Output chains", ""]
    out_blocks = list(dict.fromkeys(cell_block.get(o, o) for o in graph["outputs"]))
    seen = set()
    for ob in out_blocks:
        lines += [f"### {short(ob)}", "", "```", label(ob)]

        def walk(bid, depth):
            for src in sorted(bedges.get(bid, []), key=short):
                mark = f"  <~ {conf_of[(src, bid)]}" if (src, bid) in conf_of else ""
                if src in seen:
                    lines.append("  " * depth + f"<- {short(src)} (expanded above){mark}")
                    continue
                seen.add(src)
                lines.append("  " * depth + f"<- {label(src)}{mark}".replace("\n", "\n" + "  " * depth))
                if depth < 10:
                    walk(src, depth + 1)
        walk(ob, 1)
        lines += ["```", ""]

    lines += ["## True inputs", "", "| Block | Header | Cells | Types | Range / values | Thresholds downstream |",
              "|---|---|---|---|---|---|"]
    for b in graph["input_blocks"]:
        rng = (f"{b.get('min'):,} .. {b.get('max'):,}" if "min" in b else
               f"{b.get('min_date')} .. {b.get('max_date')}" if "min_date" in b else
               ", ".join(repr(x) for x in b.get("distinct", [])[:6]) + (" ..." if b.get("distinct_count", 0) > 6 else ""))
        hid = f" ({b['sheet_state']})" if b["sheet_state"] != "visible" else ""
        lines.append(f"| {b['range']}{hid} | {b.get('header') or ''} | {b['cells']} | {b['types']} | {rng} | "
                     f"{', '.join(f'{t:g}' for t in b['thresholds'][:8])} |")
    if graph["externals"]:
        lines += ["", f"External inputs (other workbooks, cached values only): {', '.join(graph['externals'][:20])}"]

    sections = [
        ("Dynamic references", graph["dynamic"], ["cell", "function", "formula", "base", "drivers", "candidates", "note"]),
        ("Broken references", graph["broken_refs"], ["cell", "ref", "reason", "count"]),
        ("Lookup tables with problems", [l for l in graph["lookups"] if l["problems"]],
         ["function", "table", "approximate", "problems", "count", "cells"]),
        ("Aggregates that stop short of adjacent data", graph["aggregate_gaps"],
         ["cells", "formula", "range", "excluded_after", "excluded_count", "count"]),
        ("Circular references", graph["cycles"], ["size", "cells"]),
    ]
    for title, items, keys in sections:
        lines += ["", f"## {title} ({len(items)})", ""]
        for it in items[:40]:
            lines.append("- " + _fmt(it, keys))
    if graph["outputs_secondary"]:
        lines += ["", f"## Other end-points ({len(graph['outputs_secondary'])} formula cells nothing reads, "
                      f"not on an output sheet)", "", ", ".join(graph["outputs_secondary"][:60])]
    return "\n".join(lines) + "\n"
