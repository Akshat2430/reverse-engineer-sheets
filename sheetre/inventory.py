"""Stage 1: Inventory. Read the file without calculating it.

Writes inventory.json: sheets, names, links, VBA, calc settings and flag lists.
Every flag is something a human should look at, not a verdict.
"""
import hashlib
import json
import re
import statistics
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

from . import refs as R
from .xlsxml import Package

CAP = 150                       # max items per flag list (counts are always exact)
NUM_TEXT = re.compile(r"^\s*[-+(]?\s*[$€£₹]?\s*\d{1,3}(?:,\d{2,3})+(?:\.\d+)?\s*\)?\s*$|^\s*[-+]?\d+(?:\.\d+)?\s*%?\s*$")
INDIAN = re.compile(r"^\s*[-+]?\d{1,2}(?:,\d{2})+,\d{3}(?:\.\d+)?\s*$")
TEXT_DATE = re.compile(r"^\s*(\d{1,4})[/.-](\d{1,2})[/.-](\d{1,4})\s*$")
UNIT_LABEL = re.compile(r"(?i)('000|000s|\bin thousands\b|\bthousands\b|\blakhs?\b|\bcrores?\b|\bmn\b|\bmillions?\b|\bbn\b|\busd ?k\b)")


def formula_text(v):
    if isinstance(v, ArrayFormula):
        return v.text
    return v if R.is_formula(v) else None


def luminance(color, theme_default):
    """0..1 brightness of an openpyxl Color, None if unknown."""
    if color is None:
        return theme_default
    if color.type == "rgb" and isinstance(color.rgb, str) and len(color.rgb) >= 6:
        h = color.rgb[-6:]
        r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
        return 0.2126 * r + 0.7152 * g + 0.0722 * b
    if color.type == "theme":
        base = {0: 1.0, 1: 0.0, 2: 0.9, 3: 0.25}.get(color.theme)   # lt1, dk1, lt2, dk2
        if base is None:
            return None
        tint = color.tint or 0
        return base + (1 - base) * tint if tint > 0 else base * (1 + tint)
    if color.type == "indexed":
        return {1: 1.0, 9: 1.0, 0: 0.0, 8: 0.0, 64: 0.0, 65: 1.0}.get(color.indexed)
    return None


def capped(items):
    return dict(count=len(items), items=items[:CAP], truncated=len(items) > CAP)


def run(path, out_dir):
    path, out_dir = Path(path), Path(out_dir)
    pkg = Package(path)
    wbinfo = pkg.workbook()
    wb_f = openpyxl.load_workbook(path, keep_vba=True)                  # formulas
    wb_v = openpyxl.load_workbook(path, keep_vba=True, data_only=True)  # cached values

    inv = dict(file=str(path), name=path.name, size=path.stat().st_size,
               sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
               calc=wbinfo["calc"], sheets=[], defined_names=[], external_links=wbinfo["external_links"])
    flags = defaultdict(list)
    states = {s["name"]: s["state"] for s in wbinfo["sheets"]}
    parts = {s["name"]: s["part"] for s in wbinfo["sheets"]}

    # ---------------- defined names
    for dn in wbinfo["defined_names"]:
        broken = "#REF!" in dn["refers_to"]
        inv["defined_names"].append(dict(dn, broken=broken))
        if broken:
            flags["broken_names"].append(dict(name=dn["name"], refers_to=dn["refers_to"]))
    lower = Counter(n["name"].lower().replace("_", "") for n in wbinfo["defined_names"] if not n["name"].startswith("_xl"))
    for n in wbinfo["defined_names"]:
        if lower[n["name"].lower().replace("_", "")] > 1:
            flags["similar_names"].append(n["name"])

    # ---------------- per sheet
    header_sig = {}
    col_magnitude = {}
    for ws in wb_f.worksheets:
        wv = wb_v[ws.title]
        extras = pkg.sheet_extras(parts.get(ws.title))
        info = dict(name=ws.title, state=states.get(ws.title, ws.sheet_state), dims=ws.dimensions,
                    max_row=ws.max_row, max_col=ws.max_column, formulas=0, constants=0,
                    hidden_cols=[k for k, d in ws.column_dimensions.items() if d.hidden],
                    hidden_rows=[k for k, d in ws.row_dimensions.items() if d.hidden],
                    merged=[str(m) for m in ws.merged_cells.ranges],
                    freeze=ws.freeze_panes)
        if info["state"] != "visible":
            flags["hidden_sheets"].append(dict(sheet=ws.title, state=info["state"]))
        for c in info["hidden_cols"]:
            flags["hidden_columns"].append(f"{ws.title}!{c}:{c}")
        for r in info["hidden_rows"]:
            flags["hidden_rows"].append(f"{ws.title}!{r}:{r}")
        for m in info["merged"]:
            flags["merged_ranges"].append(f"{ws.title}!{m}")

        col_vals = defaultdict(list)     # column -> [(row, value)] constants below header
        patterns = defaultdict(list)     # column -> [(row, r1c1 or None for constant)]
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if v is None:
                    continue
                key = f"{ws.title}!{cell.coordinate}"
                f = formula_text(v)
                if f:
                    info["formulas"] += 1
                    patterns[cell.column].append((cell.row, R.r1c1(f, cell.row, cell.column)))
                    _formula_flags(flags, key, f)
                    if isinstance(v, ArrayFormula):
                        flags["array_formulas"].append(dict(cell=key, ref=v.ref, formula=f))
                else:
                    info["constants"] += 1
                    patterns[cell.column].append((cell.row, None))
                    col_vals[cell.column].append((cell.row, v))
                    _value_flags(flags, key, v)
                _style_flags(flags, key, cell, v)
                if cell.comment:
                    flags["comments"].append(dict(cell=key, author=cell.comment.author, text=cell.comment.text.strip()))
        for e in extras["cached_errors"]:
            flags["cached_errors"].append(dict(cell=f"{ws.title}!{e['cell']}", error=e["error"], formula=e["has_formula"]))
        for af in extras["array_formulas"]:
            if not any(a["cell"] == f"{ws.title}!{af['cell']}" for a in flags["array_formulas"]):
                flags["array_formulas"].append(dict(cell=f"{ws.title}!{af['cell']}", ref=af["ref"], formula=af["formula"]))
        for dv in ws.data_validations.dataValidation:
            flags["data_validations"].append(dict(sheet=ws.title, sqref=str(dv.sqref), type=dv.type,
                                                  formula1=dv.formula1, formula2=dv.formula2, source="sheet"))
        for dv in extras["x14_validations"]:
            flags["data_validations"].append(dict(sheet=ws.title, sqref=dv["sqref"], type=dv["type"],
                                                  formula1=dv["formula1"], formula2=dv["formula2"], source="x14 extLst"))
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                flags["conditional_formats"].append(dict(sheet=ws.title, range=str(cf.sqref), type=rule.type,
                                                         operator=rule.operator, formula=list(rule.formula or [])))
        if extras["x14_conditional"]:
            flags["conditional_formats"].append(dict(sheet=ws.title, range="?", type="x14 extension",
                                                     operator=None, formula=[f"{extras['x14_conditional']} x14 rules (not parsed)"]))
        _pattern_flags(flags, ws.title, patterns)
        _column_type_flags(flags, ws.title, col_vals)

        # signature for sibling-sheet scale comparison: first non-empty text row
        sig = None
        for r in range(1, min(ws.max_row, 10) + 1):
            texts = tuple(str(ws.cell(r, c).value).strip() for c in range(1, min(ws.max_column, 12) + 1)
                          if isinstance(ws.cell(r, c).value, str))
            if len(texts) >= 2:
                sig = (r, texts)
                break
        if sig:
            header_sig[ws.title] = sig
            for c, vals in col_vals.items():
                nums = [abs(v) for _, v in vals if isinstance(v, (int, float)) and not isinstance(v, bool) and v]
                if len(nums) >= 5:
                    col_magnitude[(ws.title, c)] = statistics.median(nums)
        inv["sheets"].append(info)

    _sibling_scale_flags(flags, header_sig, col_magnitude)

    # ---------------- external links
    for link in wbinfo["external_links"]:
        flags["external_links"].append(dict(index=link["index"], target=link["target"],
                                            sheets=link["sheets"], cached_cells=len(link["cached"])))
    # ---------------- calc settings
    if inv["calc"].get("iterate") in ("1", "true"):
        flags["iterative_calc"].append(inv["calc"])
    if inv["calc"].get("calcMode") == "manual":
        flags["manual_calc"].append(inv["calc"])

    # ---------------- VBA
    inv["vba"] = _vba(path, out_dir)
    if inv["vba"]["modules"]:
        for m in inv["vba"]["modules"]:
            for w in m["writes"]:
                flags["vba_writes"].append(dict(module=m["name"], line=w))

    # ---------------- CF thresholds not enforced by formulas
    lits = {round(x["value"], 9) for x in flags["literals"]}
    for cf in flags["conditional_formats"]:
        for f in cf["formula"]:
            try:
                val = round(float(f), 9)
            except (TypeError, ValueError):
                continue
            if val not in lits:
                flags["cf_threshold_not_in_formulas"].append(dict(cf, threshold=val))

    for k in ("volatile", "ref_errors_in_formulas", "error_masking", "literals", "hidden_text"):
        flags[k] = group_by_column(flags[k])
    inv["flags"] = {k: capped(v) for k, v in sorted(flags.items())}
    inv["summary"] = {k: len(v) for k, v in sorted(flags.items())}
    (out_dir / "inventory.json").write_text(json.dumps(inv, indent=1, default=str))
    return inv


# ---------------------------------------------------------------- flag helpers

def _formula_flags(flags, key, f):
    funcs = {fn.name for fn in R.functions(f)}
    _, col, row = R.split_key(key)
    pat = R.r1c1(f, row, col)
    vol = funcs & R.VOLATILE
    if vol:
        flags["volatile"].append(dict(cell=key, functions=sorted(vol), formula=f, pattern=pat))
    if "#REF!" in f:
        flags["ref_errors_in_formulas"].append(dict(cell=key, formula=f, pattern=pat))
    if funcs & {"IFERROR", "IFNA", "ISERROR", "ISERR"}:
        flags["error_masking"].append(dict(cell=key, formula=f, pattern=pat))
    for r in R.refs(f, key.rsplit("!", 1)[0]):
        if r.kind == "external":
            flags["external_refs"].append(dict(cell=key, ref=r.text))
            break
    for lit in R.literals(f):
        flags["literals"].append(dict(cell=key, value=lit.value, context=lit.context, function=lit.func, formula=f, pattern=pat))


def _value_flags(flags, key, v):
    if isinstance(v, str):
        s = v
        if NUM_TEXT.match(s) or INDIAN.match(s):
            flags["text_numbers"].append(dict(cell=key, value=s, indian_grouping=bool(INDIAN.match(s))))
        m = TEXT_DATE.match(s)
        if m:
            a, b = int(m.group(1)), int(m.group(2))
            flags["text_dates"].append(dict(cell=key, value=s, ambiguous=a <= 12 and b <= 12 and len(m.group(1)) <= 2))
        if s != s.strip() or "  " in s or " " in s:
            flags["whitespace_text"].append(dict(cell=key, value=s))
        if UNIT_LABEL.search(s):
            flags["unit_labels"].append(dict(cell=key, text=s))


def _style_flags(flags, key, cell, v):
    if v in (None, ""):
        return
    font = cell.font
    fill = cell.fill
    fg = luminance(font.color if font else None, 0.0)
    if fill is None or fill.patternType in (None, "none"):
        bg = 1.0
    else:
        bg = luminance(fill.fgColor, 1.0)
    if fg is not None and bg is not None and fg > 0.9 and bg > 0.9:
        flags["hidden_text"].append(dict(cell=key, value=str(v)[:80], reason="white or near-white font on white background"))
    elif cell.number_format and ";;;" in cell.number_format:
        flags["hidden_text"].append(dict(cell=key, value=str(v)[:80], reason="number format ;;; hides the value"))


def group_by_column(items):
    """Collapse per-cell flags that share sheet, column and formula pattern (copied-down formulas)."""
    groups = {}
    for it in items:
        sheet, col, row = R.split_key(it["cell"])
        k = (sheet, col, it.get("pattern"), it.get("value") if "pattern" in it else it.get("reason"),
             it.get("context"))
        g = groups.get(k)
        if g is None:
            groups[k] = g = dict(it, cells=[], count=0)
            g.pop("cell", None)
        g["cells"].append(it["cell"])
        g["count"] += 1
    out = []
    for g in groups.values():
        rows = [R.split_key(c)[2] for c in g["cells"]]
        letter = get_column_letter(R.split_key(g["cells"][0])[1])
        sheet = g["cells"][0].rsplit("!", 1)[0]
        g["range"] = f"{sheet}!{letter}{min(rows)}" + (f":{letter}{max(rows)}" if len(rows) > 1 else "")
        g["cells"] = g["cells"][:6]
        out.append(g)
    return out


def _runs(rows):
    """Split sorted row numbers into contiguous runs."""
    run = []
    for r in rows:
        if run and r != run[-1] + 1:
            yield run
            run = []
        run.append(r)
    if run:
        yield run


def _pattern_flags(flags, sheet, patterns):
    """In each column, find formula runs where one cell differs from the majority pattern,
    and constants sitting inside a run of formulas."""
    for col, items in patterns.items():
        items.sort()
        by_row = dict(items)
        for run in _runs([r for r, _ in items]):
            if len(run) < 4:
                continue
            pats = [by_row[r] for r in run]
            formulas = [p for p in pats if p]
            if len(formulas) < 3:
                continue
            common, n = Counter(formulas).most_common(1)[0]
            if n / len(formulas) < 0.5:
                continue
            letter = get_column_letter(col)
            rng = f"{sheet}!{letter}{run[0]}:{letter}{run[-1]}"
            odd = [r for r in run if by_row[r] and by_row[r] != common]
            consts = [r for r in run if by_row[r] is None and run[0] < r]   # skip the header row
            if odd:
                flags["formula_pattern_breaks"].append(dict(range=rng, majority=common, majority_count=n,
                                                            different=[f"{letter}{r}" for r in odd][:40],
                                                            different_count=len(odd)))
            if consts and len(formulas) >= 0.5 * len(run):
                flags["constants_in_formula_runs"].append(dict(range=rng, majority=common,
                                                               constants=[f"{letter}{r}" for r in consts][:40],
                                                               constant_count=len(consts)))


def _column_type_flags(flags, sheet, col_vals):
    """Columns that are mostly text codes but contain numbers/dates: possible auto-conversion on paste."""
    import datetime as dt
    for col, vals in col_vals.items():
        body = vals[1:] if vals and isinstance(vals[0][1], str) else vals
        if len(body) < 5:
            continue
        kinds = Counter("text" if isinstance(v, str) else "date" if isinstance(v, (dt.date, dt.datetime))
                        else "number" if isinstance(v, (int, float)) and not isinstance(v, bool) else "other"
                        for _, v in body)
        if kinds["text"] / len(body) >= 0.7 and (kinds["number"] or kinds["date"]):
            letter = get_column_letter(col)
            odd = [f"{sheet}!{letter}{r}" for r, v in body if not isinstance(v, str)]
            flags["mixed_type_columns"].append(dict(column=f"{sheet}!{letter}", text=kinds["text"],
                                                    numbers=kinds["number"], dates=kinds["date"], cells=odd[:20]))


def _sibling_scale_flags(flags, header_sig, col_magnitude):
    groups = defaultdict(list)
    for sheet, sig in header_sig.items():
        groups[sig].append(sheet)
    for sig, sheets in groups.items():
        if len(sheets) < 3:
            continue
        cols = {c for (s, c) in col_magnitude if s in sheets}
        for c in cols:
            mags = {s: col_magnitude[(s, c)] for s in sheets if (s, c) in col_magnitude}
            if len(mags) < 3:
                continue
            med = statistics.median(mags.values())
            for s, m in mags.items():
                ratio = med / m if m else float("inf")
                if ratio >= 50 or ratio <= 1 / 50:
                    flags["scale_outliers"].append(dict(sheet=s, column=get_column_letter(c), median_value=m,
                                                        siblings_median=med, ratio=round(ratio, 2),
                                                        siblings=sorted(mags)))


def _vba(path, out_dir):
    out = dict(present=False, modules=[], error=None)
    try:
        from oletools.olevba import VBA_Parser
        vp = VBA_Parser(str(path))
        if not vp.detect_vba_macros():
            return out
        out["present"] = True
        vdir = out_dir / "vba"
        vdir.mkdir(parents=True, exist_ok=True)
        for _, stream, name, code in vp.extract_macros():
            lines = [l for l in code.splitlines() if not l.startswith("Attribute VB_")]
            body = "\n".join(lines).strip()
            if not body:
                continue
            procs = re.findall(r"(?im)^\s*(?:public |private )?(sub|function)\s+(\w+)", body)
            writes = [l.strip() for l in lines if re.search(r"\.(Value|Formula|Value2)\s*=", l)]
            safe = re.sub(r"[^\w.]", "_", name)
            (vdir / safe).write_text(body)
            out["modules"].append(dict(name=name, stream=stream, file=f"vba/{safe}",
                                       procedures=[f"{k.title()} {n}" for k, n in procs], writes=writes,
                                       lines=len(lines)))
    except Exception as e:                      # olevba failing is UNKNOWN, not "no macros"
        out["error"] = f"{type(e).__name__}: {e}"
    return out
