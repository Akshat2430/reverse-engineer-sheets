"""Formula parsing: references, a light function-call tree, R1C1 patterns and literals.

Built on openpyxl's tokenizer, which handles quoting, strings and nested functions.
"""
import re
from dataclasses import dataclass, field

from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.utils import column_index_from_string, get_column_letter

VOLATILE = {"TODAY", "NOW", "RAND", "RANDBETWEEN", "RANDARRAY", "OFFSET", "INDIRECT", "CELL", "INFO"}
DYNAMIC = {"OFFSET", "INDIRECT"}
COMPARE_OPS = {">", "<", ">=", "<=", "=", "<>"}
ERRORS = {"#REF!", "#N/A", "#VALUE!", "#DIV/0!", "#NUM!", "#NAME?", "#NULL!", "#SPILL!", "#CALC!", "#GETTING_DATA"}

_CELL = r"\$?[A-Z]{1,3}\$?\d+"
_COL = r"\$?[A-Z]{1,3}"
_ROW = r"\$?\d+"
RANGE_RE = re.compile(rf"^(?:({_CELL})(?::({_CELL}))?|({_COL}):({_COL})|({_ROW}):({_ROW}))$")
EXT_RE = re.compile(r"^\[(\d+)\]")


def cell_key(sheet, col, row):
    return f"{sheet}!{get_column_letter(col)}{row}"


def split_key(key):
    sheet, addr = key.rsplit("!", 1)
    m = re.match(r"([A-Z]+)(\d+)$", addr)
    return sheet, column_index_from_string(m.group(1)), int(m.group(2))


def parse_cell(a):
    m = re.match(r"^\$?([A-Z]{1,3})\$?(\d+)$", a)
    return column_index_from_string(m.group(1)), int(m.group(2))


def unquote_sheet(s):
    if s.startswith("'") and s.endswith("'"):
        return s[1:-1].replace("''", "'")
    return s


@dataclass
class Ref:
    """One reference found in a formula."""
    text: str                    # as written
    kind: str                    # cell | range | name | external | error
    sheet: str = None            # resolved sheet (None for names/errors)
    c1: int = None
    r1: int = None
    c2: int = None
    r2: int = None
    external: str = None         # external book index for kind=external
    func: str = None             # innermost enclosing function
    arg: int = None              # argument index within that function

    @property
    def whole_col(self):
        return self.r1 is None and self.c1 is not None

    @property
    def whole_row(self):
        return self.c1 is None and self.r1 is not None

    def cells(self, max_row=None, max_col=None, limit=200000):
        """Cell keys this reference covers, bounded by the sheet's used range."""
        if self.kind not in ("cell", "range"):
            return []
        c1, c2 = self.c1 or 1, self.c2 or (max_col or 1)
        r1, r2 = self.r1 or 1, self.r2 or (max_row or 1)
        if max_row:
            r2 = min(r2, max_row)
        if max_col:
            c2 = min(c2, max_col)
        n = max(0, c2 - c1 + 1) * max(0, r2 - r1 + 1)
        if n > limit:
            raise ValueError(f"range {self.text} too large ({n} cells)")
        return [cell_key(self.sheet, c, r) for c in range(c1, c2 + 1) for r in range(r1, r2 + 1)]

    @property
    def a1(self):
        if self.kind not in ("cell", "range"):
            return self.text
        if self.whole_col:
            return f"{self.sheet}!{get_column_letter(self.c1)}:{get_column_letter(self.c2)}"
        if self.whole_row:
            return f"{self.sheet}!{self.r1}:{self.r2}"
        a = f"{get_column_letter(self.c1)}{self.r1}"
        if (self.c2, self.r2) != (self.c1, self.r1):
            a += f":{get_column_letter(self.c2)}{self.r2}"
        return f"{self.sheet}!{a}"


def parse_ref(text, current_sheet):
    """Parse a RANGE operand token into a Ref (names returned as kind='name')."""
    t = text.strip()
    if t in ERRORS or t.endswith("#REF!"):
        return Ref(text, "error", sheet=current_sheet)
    ext = None
    sheet = current_sheet
    addr = t
    if "!" in t:
        sheet_part, addr = t.rsplit("!", 1)
        sheet_part = unquote_sheet(sheet_part)
        m = EXT_RE.match(sheet_part)
        if m:
            ext = m.group(1)
            sheet_part = sheet_part[m.end():]
        sheet = sheet_part
        if addr.upper().startswith("#REF"):
            return Ref(text, "error", sheet=sheet)
    m = RANGE_RE.match(addr.upper())
    if not m:
        return Ref(text, "name", sheet=None if "!" not in t else sheet, external=ext)
    a, b, ca, cb, ra, rb = m.groups()
    if a:
        c1, r1 = parse_cell(a)
        c2, r2 = parse_cell(b) if b else (c1, r1)
        r = Ref(text, "cell" if not b else "range", sheet, min(c1, c2), min(r1, r2), max(c1, c2), max(r1, r2))
    elif ca:
        i, j = column_index_from_string(ca.replace("$", "")), column_index_from_string(cb.replace("$", ""))
        r = Ref(text, "range", sheet, min(i, j), None, max(i, j), None)
    else:
        i, j = int(ra.replace("$", "")), int(rb.replace("$", ""))
        r = Ref(text, "range", sheet, None, min(i, j), None, max(i, j))
    if ext:
        r.kind, r.external = "external", ext
    return r


# ---------------------------------------------------------------- function tree

@dataclass
class Func:
    name: str
    args: list = field(default_factory=list)   # each arg: list of nodes (Token or Func)


def tokens(formula):
    if not isinstance(formula, str) or not formula.startswith("="):
        return []
    try:
        return Tokenizer(formula).items
    except Exception:
        return []


def tree(formula):
    """Nest tokens into Func nodes. Returns the top-level node list."""
    root = []
    stack = [(None, root)]                      # (Func or None, current arg list)
    for tok in tokens(formula):
        func, cur = stack[-1]
        if tok.type == Token.FUNC and tok.subtype == Token.OPEN:
            f = Func(tok.value[:-1].upper().replace("_XLFN.", "").replace("_XLWS.", ""), [[]])
            cur.append(f)
            stack.append((f, f.args[0]))
        elif tok.type == Token.FUNC and tok.subtype == Token.CLOSE:
            if len(stack) > 1:
                stack.pop()
        elif tok.type == Token.SEP and tok.subtype == Token.ARG and func is not None:
            func.args.append([])
            stack[-1] = (func, func.args[-1])
        elif tok.type == Token.WSPACE:
            continue
        else:
            cur.append(tok)
    return root


def walk(nodes, func=None, arg=None):
    """Yield (node, enclosing_func, arg_index, siblings, index) for every node."""
    for i, n in enumerate(nodes):
        yield n, func, arg, nodes, i
        if isinstance(n, Func):
            for j, a in enumerate(n.args):
                yield from walk(a, n, j)


def functions(formula):
    return [n for n, *_ in walk(tree(formula)) if isinstance(n, Func)]


def refs(formula, current_sheet):
    """All references in a formula with their enclosing function/argument."""
    out = []
    for n, func, arg, _, _ in walk(tree(formula)):
        if isinstance(n, Token) and n.type == Token.OPERAND and n.subtype == Token.RANGE:
            r = parse_ref(n.value, current_sheet)
            r.func, r.arg = (func.name if func else None), arg
            out.append(r)
    return out


def node_text(nodes):
    parts = []
    for n in nodes:
        if isinstance(n, Func):
            parts.append(f"{n.name}(" + ",".join(node_text(a) for a in n.args) + ")")
        else:
            parts.append(n.value)
    return "".join(parts)


# ---------------------------------------------------------------- literals

# (function, arg index) positions where a small integer is structural, not a business number
_STRUCTURAL = {("VLOOKUP", 2), ("VLOOKUP", 3), ("HLOOKUP", 2), ("HLOOKUP", 3), ("INDEX", 1), ("INDEX", 2),
               ("MATCH", 2), ("ROUND", 1), ("ROUNDUP", 1), ("ROUNDDOWN", 1), ("OFFSET", 1), ("OFFSET", 2),
               ("OFFSET", 3), ("OFFSET", 4), ("CHOOSE", 0), ("LEFT", 1), ("RIGHT", 1), ("MID", 1), ("MID", 2),
               ("XLOOKUP", 4), ("XLOOKUP", 5), ("WEEKDAY", 1), ("DATE", 0), ("DATE", 1), ("DATE", 2),
               ("EOMONTH", 1), ("EDATE", 1), ("LARGE", 1), ("SMALL", 1), ("TEXT", 1), ("FIXED", 1)}


@dataclass
class Literal:
    value: float
    context: str        # compare | multiply | add | arg | other
    func: str = None
    op: str = None


def literals(formula):
    """Numeric literals with a rough role. Structural ints (column indexes etc.) are skipped."""
    out = []
    for n, func, arg, sib, i in walk(tree(formula)):
        if not (isinstance(n, Token) and n.type == Token.OPERAND and n.subtype == Token.NUMBER):
            continue
        try:
            v = float(n.value)
        except ValueError:
            continue
        fname = func.name if func else None
        if fname and (fname, arg) in _STRUCTURAL:
            continue
        prev = sib[i - 1] if i > 0 else None
        nxt = sib[i + 1] if i + 1 < len(sib) else None
        ops = [t.value for t in (prev, nxt) if isinstance(t, Token) and t.type in (Token.OP_IN, Token.OP_PRE)]
        if any(o in COMPARE_OPS for o in ops):
            ctx = "compare"
        elif any(o in ("*", "/") for o in ops):
            ctx = "multiply"
        elif any(o in ("+", "-") for o in ops):
            ctx = "add"
        elif fname:
            ctx = "arg"
        else:
            ctx = "other"
        if ctx in ("add", "other", "arg") and v in (0, 1, -1):
            continue
        out.append(Literal(v, ctx, fname, next(iter(ops), None)))
    return out


def comparisons(formula):
    """(operator, literal) pairs, e.g. ('>=', 5000.0) -- candidate thresholds for probing."""
    out = []
    for n, func, arg, sib, i in walk(tree(formula)):
        if isinstance(n, Token) and n.type == Token.OP_IN and n.value in COMPARE_OPS:
            for j in (i - 1, i + 1):
                if 0 <= j < len(sib):
                    t = sib[j]
                    if isinstance(t, Token) and t.type == Token.OPERAND and t.subtype == Token.NUMBER:
                        out.append((n.value, float(t.value)))
    return out


# ---------------------------------------------------------------- R1C1 pattern

def _r1c1_cell(a, row, col):
    m = re.match(r"^(\$?)([A-Z]{1,3})(\$?)(\d+)$", a)
    cabs, cl, rabs, rw = m.groups()
    c, r = column_index_from_string(cl), int(rw)
    rp = f"R{r}" if rabs else (f"R[{r - row}]" if r != row else "R")
    cp = f"C{c}" if cabs else (f"C[{c - col}]" if c != col else "C")
    return rp + cp


def r1c1(formula, row, col):
    """Position-independent form of a formula, for spotting copy-down breaks."""
    if not isinstance(formula, str) or not formula.startswith("="):
        return None
    out = []
    for tok in tokens(formula):
        v = tok.value
        if tok.type == Token.OPERAND and tok.subtype == Token.RANGE:
            prefix, addr = (v.rsplit("!", 1) if "!" in v else ("", v))
            parts = addr.split(":")
            if all(re.match(r"^\$?[A-Z]{1,3}\$?\d+$", p.upper()) for p in parts):
                addr = ":".join(_r1c1_cell(p.upper(), row, col) for p in parts)
            v = (prefix + "!" if prefix else "") + addr
        out.append(v)
    return "=" + "".join(out)


def is_formula(v):
    return isinstance(v, str) and v.startswith("=")
