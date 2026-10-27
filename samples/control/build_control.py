"""Control workbook: a small, clean team budget with no planted problems.

    .venv/bin/python samples/control/build_control.py

The skill should return zero findings flagged as problems and a 100% proof on it.
"""
import os
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
NAME = "Team_Budget_Control.xlsx"
TEAMS = [("ENG", "Engineering", 14, 145000), ("SAL", "Sales", 9, 98000), ("MKT", "Marketing", 5, 91000),
         ("OPS", "Operations", 6, 76000), ("FIN", "Finance", 3, 104000)]

HDR = dict(fill=PatternFill("solid", fgColor="1F4E78"), font=Font(bold=True, color="FFFFFF"))


def head(ws, row, labels):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row, i, label)
        c.fill, c.font = HDR["fill"], HDR["font"]


def build(path):
    wb = openpyxl.Workbook()
    a = wb.active
    a.title = "Assumptions"
    a["A1"] = "Assumptions"; a["A1"].font = Font(bold=True, size=14)
    for r, label, val, fmt in [(3, "Benefits load (% of salary)", 0.22, "0.0%"), (4, "Travel per head (annual)", 6500, "#,##0"),
                               (5, "Software per head (annual)", 2400, "#,##0"), (6, "Merit increase", 0.04, "0.0%")]:
        a.cell(r, 1, label)
        a.cell(r, 2, val).number_format = fmt

    t = wb.create_sheet("Teams")
    head(t, 1, ["Code", "Team", "Headcount", "Avg Salary"])
    for i, (code, name, heads, sal) in enumerate(TEAMS, start=2):
        t.cell(i, 1, code); t.cell(i, 2, name); t.cell(i, 3, heads); t.cell(i, 4, sal).number_format = "#,##0"

    c = wb.create_sheet("Calc")
    head(c, 1, ["Code", "Team", "Headcount", "Salary cost", "Benefits", "Travel", "Software", "Total cost"])
    last = 1 + len(TEAMS)
    for r in range(2, last + 1):
        c[f"A{r}"] = f"=Teams!A{r}"
        c[f"B{r}"] = f"=INDEX(Teams!$B$2:$B${last},MATCH(A{r},Teams!$A$2:$A${last},0))"
        c[f"C{r}"] = f"=Teams!C{r}"
        c[f"D{r}"] = f"=C{r}*Teams!D{r}*(1+Assumptions!$B$6)"
        c[f"E{r}"] = f"=D{r}*Assumptions!$B$3"
        c[f"F{r}"] = f"=C{r}*Assumptions!$B$4"
        c[f"G{r}"] = f"=C{r}*Assumptions!$B$5"
        c[f"H{r}"] = f"=ROUND(SUM(D{r}:G{r}),2)"
        for col in "DEFGH":
            c[f"{col}{r}"].number_format = "#,##0.00"

    s = wb.create_sheet("Summary")
    s["A1"] = "FY27 Team Budget - Summary"; s["A1"].font = Font(bold=True, size=14)
    head(s, 3, ["Team", "Total cost", "Share of total"])
    for i, r in enumerate(range(4, 4 + len(TEAMS))):
        src = 2 + i
        s[f"A{r}"] = f"=Calc!B{src}"
        s[f"B{r}"] = f"=Calc!H{src}"
        s[f"C{r}"] = f"=B{r}/$B${4 + len(TEAMS) + 1}"
        s[f"B{r}"].number_format = "#,##0"
        s[f"C{r}"].number_format = "0.0%"
    tot = 4 + len(TEAMS) + 1
    s[f"A{tot}"] = "Total"; s[f"A{tot}"].font = Font(bold=True)
    s[f"B{tot}"] = f"=SUM(B4:B{tot - 2})"; s[f"B{tot}"].number_format = "#,##0"
    s[f"A{tot + 1}"] = "Total headcount"; s[f"B{tot + 1}"] = f"=SUM(Calc!C2:C{last})"
    s[f"A{tot + 2}"] = "Cost per head"; s[f"B{tot + 2}"] = f"=ROUND(B{tot}/B{tot + 1},2)"
    s[f"B{tot + 2}"].number_format = "#,##0"
    for ws, w in ((a, {"A": 30, "B": 12}), (t, {"A": 8, "B": 14, "C": 11, "D": 12}),
                  (c, {k: 13 for k in "ABCDEFGH"}), (s, {"A": 18, "B": 14, "C": 14})):
        for col, width in w.items():
            ws.column_dimensions[col].width = width
    wb.active = wb.sheetnames.index("Summary")
    wb.save(path)
    # give formula cells a numeric cached value so Excel loads them cleanly
    tmp = str(path) + ".tmp"
    with zipfile.ZipFile(path) as zi, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename.startswith("xl/worksheets/sheet"):
                data = data.replace(b"<v></v>", b"<v>0</v>")
            zo.writestr(it, data)
    os.replace(tmp, path)


def excel_pass(src, dst):
    import xlwings as xw
    app = xw.apps.active if xw.apps.count else xw.App(visible=False)
    app.display_alerts = False
    b = app.books.open(str(src), update_links=False)
    try:
        app.calculate()
        b.save(str(dst))
    finally:
        b.close()


if __name__ == "__main__":
    pre = REPO / "runs" / "_build" / ("pre_" + NAME)
    pre.parent.mkdir(parents=True, exist_ok=True)
    build(pre)
    excel_pass(pre, HERE / NAME)
    pre.unlink()
    print(HERE / NAME)
