"""Inventory and map on a small synthetic workbook (no Excel needed)."""
import json

import openpyxl
import pytest
from openpyxl.styles import Font
from openpyxl.workbook.defined_name import DefinedName

from sheetre import inventory, mapgraph, report


@pytest.fixture
def book(tmp_path):
    wb = openpyxl.Workbook()
    d = wb.active
    d.title = "Data"
    d.append(["Code", "Qty", "Price"])
    rows = [("A1x", 3, 10), ("B2x", 5, 12), ("C3x", "7", 9), ("D4x", 2, 11), ("E5x", 4, 8), ("F6x", 6, 7)]
    for r in rows:
        d.append(r)
    d.append(["G7x", 1, 5])                       # row 8: data continues past the SUM range
    d["E2"] = "secret"
    d["E2"].font = Font(color="FFFFFF")
    t = wb.create_sheet("Lookup")
    for i, (k, v) in enumerate([(0, 0.0), (10, 0.05), (30, 0.1), (20, 0.08)], start=1):   # unsorted
        t.cell(i, 1, k); t.cell(i, 2, v)
    m = wb.create_sheet("Model")
    for r in range(2, 9):
        m[f"A{r}"] = f"=Data!B{r}*Data!C{r}"
        m[f"B{r}"] = f"=VLOOKUP(A{r},Lookup!$A$1:$B$4,2,TRUE)"
    m["A5"] = 99                                  # typed value inside a formula run
    m["B6"] = "=VLOOKUP(A6,Lookup!$A$1:$B$4,2,TRUE)*1.2"   # pattern break
    m["C1"] = "=C2*0.1"
    m["C2"] = "=C1+100"                           # circular
    h = wb.create_sheet("Hidden")
    h["A1"] = 5
    h.sheet_state = "veryHidden"
    s = wb.create_sheet("Summary")
    s["A1"] = "=SUM(Model!A2:A7)+Hidden!A1"       # stops short of row 8
    s["A2"] = "=IFERROR(SUMPRODUCT(Model!A2:A8,Model!B2:B8),0)"
    s["A3"] = "=OFFSET(Model!A2,Data!B3,0)+GoodName+BadName"
    s["A4"] = "=C2+Model!C2"
    wb.defined_names["GoodName"] = DefinedName("GoodName", attr_text="Data!$C$2")
    wb.defined_names["BadName"] = DefinedName("BadName", attr_text="Data!#REF!")
    p = tmp_path / "synthetic.xlsx"
    wb.save(p)
    return p


def test_inventory_flags(book, tmp_path):
    inv = inventory.run(book, tmp_path)
    s = inv["summary"]
    assert s["hidden_sheets"] == 1 and inv["flags"]["hidden_sheets"]["items"][0]["state"] == "veryHidden"
    assert s["hidden_text"] == 1
    assert s["text_numbers"] == 1
    assert s["broken_names"] == 1
    assert s["volatile"] >= 1                     # OFFSET
    assert s["error_masking"] == 1
    assert s["formula_pattern_breaks"] >= 1
    assert s["constants_in_formula_runs"] >= 1
    lit = [g for g in inv["flags"]["literals"]["items"] if g["value"] == 1.2]
    assert lit and lit[0]["context"] == "multiply"


def test_map_graph(book, tmp_path):
    g = mapgraph.build(book, tmp_path)
    assert "Summary!A1" in g["outputs"]
    assert "Hidden!A1" in g["inputs"] and "Data!C4" in g["inputs"]
    unsorted = [l for l in g["lookups"] if any("NOT sorted" in p for p in l["problems"])]
    assert unsorted
    assert any(x["excluded_count"] == 1 for x in g["aggregate_gaps"])
    assert g["cycles"]
    assert any(d["function"] == "OFFSET" for d in g["dynamic"])
    assert any("BadName" in b["ref"] for b in g["broken_refs"])
    md = (tmp_path / "map.md").read_text()
    assert "typed value(s) inside this formula run" in md and "use a different formula" in md


def test_ledger_checker(tmp_path):
    cols = "id,location,finding,category,tag,problem,evidence,settle_by,ask\n"
    (tmp_path / "ledger.csv").write_text(cols + 'L01,Model!B2,Rate lookup is approximate,lookup,SEEN,Y,"Model!B2 =VLOOKUP(...,TRUE)",,\n'
                                         'L02,Model!B6,The 1.2 is a premium,hardcoded_literal,INFERRED,Y,Model!B6 *1.2,owner confirms,\n')
    assert report.check_ledger(tmp_path) == []
    assert (tmp_path / "ledger.md").exists()
    (tmp_path / "ledger.csv").write_text(cols + "L01,X,thing,c,SEEN,Y,I think so,,\nL01,Y,other,c,MAYBE,Q,,,\n"
                                         "L03,Z,why,c,UNKNOWN,N,inventory,,\n")
    probs = report.check_ledger(tmp_path)
    assert any("SEEN needs evidence" in p for p in probs)
    assert any("duplicate id" in p for p in probs)
    assert any("tag must be" in p for p in probs)
    assert any("UNKNOWN needs" in p for p in probs)


def test_report_html_renders(tmp_path):
    html = report.to_html("# T\n\n| a | b |\n|---|---|\n| 1 | x \\| y |\n\n- one\n- two\n\n```\ncode\n```\n", "T")
    assert "<table>" in html and "<li>one</li>" in html and "<pre>code</pre>" in html and "prefers-color-scheme" in html
